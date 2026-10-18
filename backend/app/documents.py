import ast
import io
import re
import zipfile
from decimal import Decimal, localcontext
from pathlib import Path
from PIL import Image
import pymupdf
from docx import Document
from docx.shared import Inches as DocInches, Pt as DocPt, RGBColor as DocColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from openpyxl import Workbook,load_workbook
from openpyxl.chart import BarChart,LineChart,PieChart,Reference
from openpyxl.styles import Alignment,Border,Font,PatternFill,Side
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field
from .google import IntegrationError

class Section(BaseModel):
    heading:str
    paragraphs:list[str]=Field(default_factory=list)
    bullets:list[str]=Field(default_factory=list)
    table_headers:list[str]=Field(default_factory=list,max_length=12)
    table_rows:list[list[str]]=Field(default_factory=list,max_length=200)
    source_refs:list[str]=Field(default_factory=list,max_length=30)
class Slide(BaseModel):
    title:str
    bullets:list[str]=Field(max_length=6)
    notes:str=''
    source_refs:list[str]=Field(default_factory=list,max_length=30)
class Calculation(BaseModel):
    label:str
    expression:str
    source_refs:list[str]=Field(default_factory=list,max_length=30)
class WorkbookChart(BaseModel):
    kind:str='bar'
    title:str=''
    categories_column:int=1
    data_columns:list[int]=Field(default_factory=list,max_length=12)
class WorkbookSheet(BaseModel):
    name:str
    layout:str='standard'
    headers:list[str]=Field(default_factory=list,max_length=40)
    rows:list[list[str]]=Field(default_factory=list,max_length=500)
    notes:list[str]=Field(default_factory=list,max_length=30)
    source_refs:list[str]=Field(default_factory=list,max_length=30)
    chart:WorkbookChart|None=None
class Draft(BaseModel):
    title:str
    summary:str
    needs_attention:bool=False
    formats:list[str]=Field(default_factory=lambda:['docx'])
    sections:list[Section]=Field(default_factory=list,max_length=60)
    slides:list[Slide]=Field(default_factory=list,max_length=40)
    sources:list[str]=Field(default_factory=list)
    assumptions:list[str]=Field(default_factory=list)
    review_notes:list[str]=Field(default_factory=list)
    calculations:list[Calculation]=Field(default_factory=list,max_length=50)
    spreadsheets:list[WorkbookSheet]=Field(default_factory=list,max_length=20)

def extract(name,mime,data):
    if mime=='application/pdf' or name.lower().endswith('.pdf'):
        with pymupdf.open(stream=data,filetype='pdf') as pdf:
            if len(pdf)>120: raise IntegrationError(f'{name}: PDF exceeds the 120 page limit.')
            pages=[]
            for i,page in enumerate(pdf):
                text=page.get_text().strip()
                if len(text)<20:
                    raise IntegrationError(f'{name}: page {i+1} has little or no extractable text. OCR is needed before this assignment can be processed.')
                pages.append(f'[Page {i+1}]\n{text}')
            return '\n'.join(pages)
    if mime.endswith('wordprocessingml.document') or mime.endswith('presentationml.presentation') or name.lower().endswith(('.docx','.pptx')):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(x.file_size for x in archive.infolist())>60*1024*1024:
                raise IntegrationError(f'{name}: expanded document is too large.')
        if mime.endswith('wordprocessingml.document') or name.lower().endswith('.docx'):
            doc=Document(io.BytesIO(data))
            return '\n'.join([p.text for p in doc.paragraphs]+[' | '.join(cell.text for cell in row.cells) for t in doc.tables for row in t.rows])
        pres=Presentation(io.BytesIO(data))
        return '\n'.join(f'[Slide {i+1}]\n'+'\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame) for i,slide in enumerate(pres.slides))
    if mime.endswith('spreadsheetml.sheet') or name.lower().endswith('.xlsx'):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist())>60*1024*1024:
                raise IntegrationError(f'{name}: expanded workbook is too large.')
        workbook=load_workbook(io.BytesIO(data),read_only=True,data_only=False)
        sheets=[]
        for worksheet in workbook.worksheets:
            lines=[]
            for row in worksheet.iter_rows(values_only=True):
                values=[str(value) if value is not None else '' for value in row]
                while values and not values[-1]: values.pop()
                if any(values): lines.append(' | '.join(values))
            if lines: sheets.append(f'[Sheet: {worksheet.title}]\n'+'\n'.join(lines))
        workbook.close()
        return '\n\n'.join(sheets)
    if mime.startswith('text/') or name.lower().endswith(('.txt','.csv')): return data.decode('utf-8',errors='replace')
    raise IntegrationError(f'Unsupported attachment: {name}. Supported: PDF, DOCX, PPTX, XLSX, TXT, CSV, PNG, JPEG, WebP and GIF.')

def extract_pdf_with_images(name,data):
    """Extract selectable PDF text and render low-text pages for Gemini vision."""
    with pymupdf.open(stream=data,filetype='pdf') as pdf:
        if len(pdf)>120: raise IntegrationError(f'{name}: PDF exceeds the 120 page limit.')
        pages=[]; images=[]
        for index,page in enumerate(pdf):
            text=page.get_text().strip()
            if len(text)<20:
                scale=min(2.0,2400/max(page.rect.width,page.rect.height,1))
                try:
                    pix=page.get_pixmap(matrix=pymupdf.Matrix(scale,scale),colorspace=pymupdf.csRGB,alpha=False)
                    image=Image.frombytes('RGB',(pix.width,pix.height),pix.samples)
                    image.thumbnail((1800,2400),Image.Resampling.LANCZOS)
                    output=io.BytesIO(); image.save(output,format='JPEG',quality=88,optimize=True)
                except Exception as error:
                    raise IntegrationError(f'{name}: could not render scanned page {index+1} ({error}).')
                images.append({'filename':f'{name} page {index+1}','mime_type':'image/jpeg','data':output.getvalue(),'include_in_output':False})
                pages.append(f'[Page {index+1}: little selectable text; page image supplied to Gemini for visual reading]')
            else:
                pages.append(f'[Page {index+1}]\n{text}')
        return '\n'.join(pages),images

def extract_pdf_tables(data):
    """Extract selectable PDF tables as records for data-driven deliverables."""
    tables=[]
    with pymupdf.open(stream=data,filetype='pdf') as pdf:
        for page in pdf:
            try: found=page.find_tables().tables
            except Exception: continue
            for table in found:
                values=table.extract()
                if not values or len(values)<2 or not values[0]: continue
                headers=[re.sub(r'[^a-z0-9]+','_',str(value or '').strip().lower()).strip('_') for value in values[0]]
                if not any(headers): continue
                rows=[]
                for values_row in values[1:]:
                    cells=[re.sub(r'\s+',' ',str(value or '')).strip() for value in values_row]
                    if not any(cells): continue
                    rows.append({header:cells[i] if i<len(cells) else '' for i,header in enumerate(headers) if header})
                if rows: tables.append({'headers':headers,'rows':rows})
    return tables

def parse_name_card_request(instructions):
    """Read explicit name-plate/name-badge measurements and typography from instructions."""
    text=' '.join(instructions.split())
    lower=text.lower()
    if 'name plate' not in lower or 'name badge' not in lower or '.docx' not in lower: return None
    def block(start_marker,end_marker=None):
        start=lower.find(start_marker)
        if start<0: return ''
        end=lower.find(end_marker,start+len(start_marker)) if end_marker else -1
        return text[start:end if end>=0 else None]
    plate=block('formatting details for name plate','formatting details for name badges')
    badge=block('formatting details for name badges')
    def number(section,label):
        match=re.search(r'\b'+re.escape(label)+r'\s*:\s*(\d+(?:\.\d+)?)',section,re.I)
        return float(match.group(1)) if match else None
    def font(section,label,default):
        match=re.search(r'\b'+re.escape(label)+r'\s*:\s*([^\n;]+)',section,re.I)
        return re.sub(r'\s*\([^)]*\)','',match.group(1)).strip() if match else default
    spec={
        'plate_width':number(plate,'Cell Width'),'plate_height':number(plate,'Cell Height'),
        'badge_width':number(badge,'Badge Width'),'badge_height':number(badge,'Badge Height'),
        'plate_name_size':number(plate,'Name Size'),'plate_org_size':number(plate,'Organization Size'),
        'badge_name_size':number(badge,'Name Font Size'),'badge_org_size':number(badge,'Organization Font Size'),
        'plate_name_font':font(plate,'Name Font','Verdana'),'plate_org_font':font(plate,'Organization Font','Verdana'),
        'badge_name_font':font(badge,'Name Font','Verdana'),'badge_org_font':font(badge,'Organization Font','Verdana'),
        'plate_name_bold':'bold' in plate.lower(),'badge_name_bold':'bold' in badge.lower(),
    }
    required=('plate_width','plate_height','badge_width','badge_height','plate_name_size','plate_org_size','badge_name_size','badge_org_size')
    if any(spec[key] is None for key in required): return None
    return spec

def merge_table_records(tables):
    """Join matching PDF tables by row order when one table holds names and another holds details."""
    primary=next((table for table in tables if any('first_name' in row and 'organization' in row for row in table['rows'])),None)
    if not primary: return []
    records=[dict(row) for row in primary['rows']]
    for table in tables:
        if table is primary or len(table['rows'])!=len(records): continue
        for index,row in enumerate(table['rows']):
            for key,value in row.items():
                if value and key not in records[index]: records[index][key]=value
    return [row for row in records if row.get('first_name') and row.get('last_name') and row.get('organization')]

def _set_cell_borders(cell,visible=True):
    tc_pr=cell._tc.get_or_add_tcPr(); borders=tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders=OxmlElement('w:tcBorders'); tc_pr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag='w:'+edge; element=borders.find(qn(tag))
        if element is None: element=OxmlElement(tag); borders.append(element)
        element.set(qn('w:val'),'single' if visible else 'nil')
        if visible:
            element.set(qn('w:sz'),'6'); element.set(qn('w:space'),'0'); element.set(qn('w:color'),'9EAAA0')

def _set_cell_margins(cell,top=0,start=0,bottom=0,end=0):
    tc=cell._tc; tc_pr=tc.get_or_add_tcPr(); margins=tc_pr.first_child_found_in('w:tcMar')
    if margins is None: margins=OxmlElement('w:tcMar'); tc_pr.append(margins)
    for edge,value in (('top',top),('start',start),('bottom',bottom),('end',end)):
        node=margins.find(qn('w:'+edge))
        if node is None: node=OxmlElement('w:'+edge); margins.append(node)
        node.set(qn('w:w'),str(round(value*1440))); node.set(qn('w:type'),'dxa')

def _fixed_table(table,widths):
    table.autofit=False; table.alignment=WD_TABLE_ALIGNMENT.CENTER
    layout=OxmlElement('w:tblLayout'); layout.set(qn('w:type'),'fixed'); table._tbl.tblPr.append(layout)
    for row in table.rows:
        for cell,width in zip(row.cells,widths): cell.width=DocInches(width)

def _set_label_run(run,font,size,bold=False):
    run.font.name=font; run.font.size=DocPt(size); run.font.bold=bold
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),font)

def _page_setup(doc,vertical_margin):
    section=doc.sections[0]
    section.page_width=DocInches(8.5); section.page_height=DocInches(11)
    section.top_margin=DocInches(vertical_margin); section.bottom_margin=DocInches(vertical_margin)
    section.left_margin=DocInches(.25); section.right_margin=DocInches(.25)

def render_name_card_files(records,spec,images,folder):
    """Create separate print-ready DOCX files from a tabular source and rubric."""
    assets={image['filename'].lower():image for image in prepare_images(images)}
    ims=next((asset for name,asset in assets.items() if 'ims' in name),None)
    hec=next((asset for name,asset in assets.items() if 'hec' in name),None)
    output=[]

    plates=Document(); _page_setup(plates,.5)
    for start in range(0,len(records),6):
        if start: plates.add_page_break()
        table=plates.add_table(rows=0,cols=1); _fixed_table(table,[spec['plate_width']])
        for record in records[start:start+6]:
            cell=table.add_row().cells[0]; cell.width=DocInches(spec['plate_width'])
            table.rows[-1].height=DocInches(spec['plate_height']); table.rows[-1].height_rule=WD_ROW_HEIGHT_RULE.EXACTLY
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; _set_cell_borders(cell); _set_cell_margins(cell,0,.08,0,.08)
            cell.text=''; p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=DocPt(2)
            r=p.add_run((record.get('first_name','')+' '+record.get('last_name','')).strip())
            _set_label_run(r,spec['plate_name_font'],spec['plate_name_size'],spec['plate_name_bold'])
            p=cell.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=DocPt(0); p.paragraph_format.space_after=DocPt(0)
            r=p.add_run(record['organization']); _set_label_run(r,spec['plate_org_font'],spec['plate_org_size'])
    path=folder/'Name Plates.docx'; plates.save(path); output.append({'name':path.name,'kind':'DOCX'})

    badges=Document(); _page_setup(badges,.25)
    def set_logo(cell,asset,align):
        if asset is None: return
        paragraph=cell.paragraphs[0]; paragraph.alignment=align
        paragraph.paragraph_format.space_after=DocPt(0); paragraph.paragraph_format.space_before=DocPt(0)
        paragraph.add_run().add_picture(io.BytesIO(asset['data']),height=DocInches(.38))
    for start in range(0,len(records),8):
        if start: badges.add_page_break()
        table=badges.add_table(rows=4,cols=2); _fixed_table(table,[spec['badge_width'],spec['badge_width']])
        for slot in range(8):
            cell=table.cell(slot//2,slot%2); row_index=slot//2
            if slot%2==0: table.rows[row_index].height=DocInches(spec['badge_height']); table.rows[row_index].height_rule=WD_ROW_HEIGHT_RULE.EXACTLY
            cell.width=DocInches(spec['badge_width']); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            record_index=start+slot
            visible=record_index<len(records); _set_cell_borders(cell,visible); _set_cell_margins(cell,.04,.08,.04,.08)
            if not visible: cell.text=''; continue
            cell.text=''; p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=DocPt(0)
            logos=cell.add_table(rows=1,cols=2); _fixed_table(logos,[spec['badge_width']/2-.08,spec['badge_width']/2-.08])
            set_logo(logos.cell(0,0),ims,WD_ALIGN_PARAGRAPH.LEFT); set_logo(logos.cell(0,1),hec,WD_ALIGN_PARAGRAPH.RIGHT)
            for logo_cell in logos.rows[0].cells: _set_cell_borders(logo_cell,False); _set_cell_margins(logo_cell,0,0,0,0)
            p=cell.add_paragraph('Training on '+record.get('training','')); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=DocPt(2); p.paragraph_format.space_after=DocPt(1)
            for run in p.runs: _set_label_run(run,spec['badge_org_font'],spec['badge_org_size'])
            p=cell.add_paragraph((record.get('first_name','')+' '+record.get('last_name','')).strip()); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=DocPt(1); p.paragraph_format.space_after=DocPt(0)
            for run in p.runs: _set_label_run(run,spec['badge_name_font'],spec['badge_name_size'],spec['badge_name_bold'])
            p=cell.add_paragraph(record['organization']); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=DocPt(0); p.paragraph_format.space_after=DocPt(0)
            for run in p.runs: _set_label_run(run,spec['badge_org_font'],spec['badge_org_size'])
    path=folder/'Name Badges.docx'; badges.save(path); output.append({'name':path.name,'kind':'DOCX'})
    return output

def calculate(expression):
    """Only arithmetic; never execute model code, functions, names or attributes."""
    if len(expression)>200: raise ValueError('Expression too long')
    tree=ast.parse(expression,mode='eval')
    def visit(node):
        if isinstance(node,ast.Constant) and type(node.value) in (int,float): return Decimal(str(node.value))
        if isinstance(node,ast.UnaryOp) and isinstance(node.op,(ast.UAdd,ast.USub)):
            return visit(node.operand)*(-1 if isinstance(node.op,ast.USub) else 1)
        if isinstance(node,ast.BinOp):
            a,b=visit(node.left),visit(node.right)
            if isinstance(node.op,ast.Add): return a+b
            if isinstance(node.op,ast.Sub): return a-b
            if isinstance(node.op,ast.Mult): return a*b
            if isinstance(node.op,ast.Div): return a/b
            if isinstance(node.op,ast.Pow) and b==int(b) and abs(b)<=100: return a**int(b)
        raise ValueError('Only numeric arithmetic is supported')
    with localcontext() as context:
        context.prec=28
        value=visit(tree.body)
        if not value.is_finite() or abs(value)>Decimal('1e40'): raise ValueError('Result exceeds supported range')
        return format(value,'.10f').rstrip('0').rstrip('.') or '0'

def calculation_notes(draft):
    notes=[]
    for item in draft.calculations[:50]:
        try: note=f'{item.label}: {item.expression} = {calculate(item.expression)}'
        except Exception: note=f'{item.label}: could not evaluate {item.expression}'
        if item.source_refs: note+=' (Sources: '+('; '.join(item.source_refs))+')'
        notes.append(note)
    return notes

IMAGE_FORMATS={'image/png':'PNG','image/jpeg':'JPEG','image/jpg':'JPEG','image/webp':'WEBP','image/gif':'GIF'}

def validate_image(name,mime,data):
    if mime not in IMAGE_FORMATS:
        raise IntegrationError(f'{name}: unsupported image type. Supported: PNG, JPEG, WebP and GIF.')
    try:
        with Image.open(io.BytesIO(data)) as image:
            if image.format!=IMAGE_FORMATS[mime]: raise ValueError('Image content does not match its file type.')
            if image.width*image.height>40_000_000: raise ValueError('Image dimensions exceed 40 megapixels.')
            image.verify()
    except Exception as error:
        raise IntegrationError(f'{name}: image could not be read ({error}).')

def prepare_images(assets):
    prepared=[]
    for asset in assets:
        try:
            with Image.open(io.BytesIO(asset['data'])) as source:
                source.seek(0)
                image=source.convert('RGBA' if 'A' in source.getbands() or 'transparency' in source.info else 'RGB')
                image.thumbnail((2400,2400),Image.Resampling.LANCZOS)
                output=io.BytesIO(); image.save(output,format='PNG')
                prepared.append({'filename':asset['filename'],'data':output.getvalue(),'size':image.size})
        except Exception as error:
            raise IntegrationError(f"{asset['filename']}: could not prepare image for the document ({error}).")
    return prepared

SAFE_EXCEL_FUNCTIONS={
    'ABS','AND','AVERAGE','AVERAGEIF','AVERAGEIFS','COUNT','COUNTA','COUNTIF','COUNTIFS',
    'CEILING','FLOOR','IF','IFERROR','IFS','INDEX','INT','MAX','MATCH','MIN','MOD','NOT','OR',
    'PRODUCT','ROUND','ROUNDDOWN','ROUNDUP','SQRT','STDEV','SUBTOTAL','SUM','SUMIF','SUMIFS',
    'SUMPRODUCT','VLOOKUP','XLOOKUP'
}
FORMULA_CHARS=re.compile(r'=[A-Za-z0-9_.$:!,"\s()+*/^%<>=&\-]+\Z')

def _safe_excel_formula(value):
    if len(value)>500 or not FORMULA_CHARS.fullmatch(value): return False
    functions=re.findall(r'([A-Za-z_][A-Za-z0-9_.]*)\s*\(',value)
    return all(name.upper() in SAFE_EXCEL_FUNCTIONS for name in functions)

def _number_value(value):
    text=str(value).strip()
    if not text or text.startswith('=') or text.startswith("'"): return value
    percent=text.endswith('%')
    number=text[:-1] if percent else text
    if not re.fullmatch(r'-?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?',number): return value
    try:
        parsed=float(number.replace(',',''))
        return parsed/100 if percent else parsed
    except ValueError: return value

def _unique_sheet_title(title,used):
    title=re.sub(r'[\\/*?:\[\]]',' ',title).strip()[:31] or 'Sheet'
    candidate=title; index=2
    while candidate.casefold() in used:
        suffix=f' {index}'; candidate=title[:31-len(suffix)]+suffix; index+=1
    used.add(candidate.casefold())
    return candidate

def _accounting_layout(sheet):
    layout=sheet.layout.lower().replace('-','_').replace(' ','_')
    if layout not in {'standard','accounting','income_statement','balance_sheet','cash_flow','budget','ratio_analysis'}:
        layout='standard'
    if layout=='standard':
        name=sheet.name.lower()
        for candidate,words in (
            ('income_statement',('income statement','profit and loss','profit & loss')),
            ('balance_sheet',('balance sheet','statement of financial position')),
            ('cash_flow',('cash flow','cashflow')),
            ('budget',('budget','forecast')),
            ('ratio_analysis',('ratio','financial analysis')),
        ):
            if any(word in name for word in words): return candidate
    return layout

def _is_total_row(label):
    value=label.strip().lower()
    return bool(re.match(r'^(total\b|net\b|gross profit\b|operating income\b|operating profit\b|ending (cash|balance)\b|cash flow from\b)',value))

def _make_workbook(draft,folder):
    if not draft.spreadsheets: raise IntegrationError('Gemini requested an XLSX file but returned no spreadsheet sheets.')
    workbook=Workbook(); overview=workbook.active; overview.title='Overview'; overview.sheet_view.showGridLines=False
    overview['A1']=draft.title; overview['A1'].font=Font(name='Calibri',size=20,bold=True,color='FFFFFF')
    overview['A1'].fill=PatternFill('solid',fgColor='23493A'); overview.merge_cells(start_row=1,start_column=1,end_row=1,end_column=2)
    overview['A3']='Assignment summary'; overview['B3']=draft.summary
    row=5
    for title,items in (('Assumptions',draft.assumptions),('Review notes',draft.review_notes),('Sources supplied',draft.sources)):
        if not items: continue
        overview.cell(row,1,title).font=Font(bold=True,color='23493A'); row+=1
        for item in items:
            overview.cell(row,1,'• '+str(item)); overview.merge_cells(start_row=row,start_column=1,end_row=row,end_column=2); row+=1
        row+=1
    overview.column_dimensions['A'].width=28; overview.column_dimensions['B'].width=88
    for cells in overview.iter_rows():
        for cell in cells: cell.alignment=Alignment(vertical='top',wrap_text=True)
    overview.freeze_panes='A3'; overview.sheet_properties.pageSetUpPr.fitToPage=True
    used={'overview'}
    source_notes=[]
    for section in draft.sections: source_notes.extend(section.source_refs)
    for slide in draft.slides: source_notes.extend(slide.source_refs)
    for calculation in draft.calculations: source_notes.extend(calculation.source_refs)
    for sheet in draft.spreadsheets[:20]:
        title=_unique_sheet_title(sheet.name,used)
        ws=workbook.create_sheet(title); ws.sheet_view.showGridLines=False; ws.sheet_properties.tabColor='78985E'
        headers=sheet.headers or [f'Column {index+1}' for index in range(max((len(r) for r in sheet.rows),default=1))]
        ws.append(headers)
        for cell in ws[1]:
            cell.font=Font(name='Calibri',bold=True,color='FFFFFF')
            cell.fill=PatternFill('solid',fgColor='23493A')
            cell.alignment=Alignment(vertical='center',wrap_text=True)
            cell.border=Border(bottom=Side(style='medium',color='C2E49A'))
        ws.row_dimensions[1].height=30; ws.freeze_panes='A2'
        layout=_accounting_layout(sheet)
        financial=layout in {'accounting','income_statement','balance_sheet','cash_flow','budget'}
        max_col=max(1,len(headers)); max_row=1
        currency_words=('amount','revenue','sales','expense','cost','profit','assets','liabilities','equity','cash','balance','variance','dollars','debit','credit')
        ratio_words=('percent','percentage','margin','return','growth rate')
        for row_index,values in enumerate(sheet.rows[:500],start=2):
            rendered=[]
            for col_index,header in enumerate(headers,start=1):
                raw=str(values[col_index-1]) if col_index<=len(values) else ''
                cell=ws.cell(row_index,col_index)
                if raw.startswith('='):
                    if _safe_excel_formula(raw): cell.value=raw
                    else:
                        cell.value=raw; cell.data_type='s'; draft.needs_attention=True
                        note=f"Unsafe formula in {sheet.name}, {get_column_letter(col_index)}{row_index} was stored as text."
                        if note not in draft.review_notes: draft.review_notes.append(note)
                else:
                    numeric=_number_value(raw); cell.value=numeric
                cell.alignment=Alignment(vertical='top',wrap_text=True)
                normalized=header.lower()
                if financial and any(word in normalized for word in currency_words): cell.number_format='#,##0.00;[Red](#,##0.00);-'
                if any(word in normalized for word in ratio_words): cell.number_format='0.0%'
                if layout=='ratio_analysis' and 'ratio' in normalized: cell.number_format='0.00x'
                rendered.append(cell)
            label=str(values[0] if values else '').strip()
            if _is_total_row(label):
                for cell in rendered:
                    cell.font=Font(name='Calibri',bold=True,color='23493A')
                    cell.border=Border(top=Side(style='thin',color='78985E'))
            max_col=max(max_col,len(headers)); max_row=row_index
        if max_row>1: ws.auto_filter.ref=f'A1:{get_column_letter(max_col)}{max_row}'
        for col_index,header in enumerate(headers,start=1):
            values=[str(header)]+[str(row[col_index-1]) for row in sheet.rows[:500] if col_index<=len(row)]
            width=min(42,max(12,max((len(line) for value in values for line in value.splitlines()),default=12)+2))
            ws.column_dimensions[get_column_letter(col_index)].width=width
        ws.page_setup.orientation='landscape'; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
        ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.print_title_rows='1:1'
        if sheet.notes or sheet.source_refs:
            notes_row=max_row+2; ws.cell(notes_row,1,'Notes and sources').font=Font(bold=True,color='23493A')
            for note in sheet.notes: notes_row+=1; ws.cell(notes_row,1,str(note))
            for ref in sheet.source_refs:
                notes_row+=1; ws.cell(notes_row,1,'Source: '+str(ref)); source_notes.append(str(ref))
        else: source_notes.extend(sheet.source_refs)
        chart=sheet.chart
        if chart:
            kind=chart.kind.lower()
            chart_class={'bar':BarChart,'column':BarChart,'line':LineChart,'pie':PieChart}.get(kind)
            if kind=='pie' and len(chart.data_columns)!=1: chart_class=None
            last_data_row=1+min(len(sheet.rows),500)
            if chart_class and last_data_row>1 and chart.categories_column>0 and chart.categories_column<=max_col and chart.data_columns:
                chart_obj=chart_class(); chart_obj.title=chart.title or sheet.name
                chart_obj.height=8; chart_obj.width=16
                if kind in ('bar','column'):
                    chart_obj.type='col' if kind=='column' else 'bar'; chart_obj.style=10; chart_obj.y_axis.title='Value'; chart_obj.x_axis.title=headers[chart.categories_column-1]
                for data_col in chart.data_columns:
                    if 1<=data_col<=max_col:
                        chart_obj.add_data(Reference(ws,min_col=data_col,max_col=data_col,min_row=1,max_row=last_data_row),titles_from_data=True)
                chart_obj.set_categories(Reference(ws,min_col=chart.categories_column,max_col=chart.categories_column,min_row=2,max_row=last_data_row))
                if chart_obj.series: ws.add_chart(chart_obj,f'{get_column_letter(max_col+2)}2')
            else:
                draft.needs_attention=True
                draft.review_notes.append(f"The requested chart on sheet '{sheet.name}' could not be built from the supplied columns.")
    unique_refs=list(dict.fromkeys(source_notes))
    if unique_refs:
        sources_ws=workbook.create_sheet(_unique_sheet_title('Source Notes',used)); sources_ws.append(['Source reference'])
        for cell in sources_ws[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='23493A')
        for reference in unique_refs: sources_ws.append([reference])
        sources_ws.column_dimensions['A'].width=80; sources_ws.freeze_panes='A2'
    path=folder/'assignment.xlsx'; workbook.save(path)
    return {'name':path.name,'kind':'XLSX'}

def render(draft:Draft,folder:Path,image_assets=None):
    folder.mkdir(parents=True,exist_ok=True)
    images=prepare_images(image_assets or [])
    files=[]
    formats=set(draft.formats)
    if not formats or formats-{'docx','pptx','xlsx'}: raise IntegrationError('Assignment requests an output format this version cannot create. Supported outputs are DOCX, PPTX and XLSX.')
    if 'xlsx' in formats and not draft.spreadsheets:
        if formats-{'xlsx'}:
            draft.needs_attention=True
            draft.review_notes.append('The assignment requests an XLSX workbook, but the model did not provide any workbook sheets. The available document is only a supporting draft.')
            formats.remove('xlsx')
        else: raise IntegrationError('Gemini requested an XLSX file but returned no spreadsheet sheets.')
    if 'docx' in formats and not draft.sections: raise IntegrationError('The model returned no document sections. Retry generation.')
    if 'pptx' in formats and not draft.slides: raise IntegrationError('The model returned no slides. Retry generation.')
    calculations=calculation_notes(draft)
    if 'docx' in formats:
        if not draft.sections: raise IntegrationError('The model returned no document sections. Retry generation.')
        doc=Document()
        doc.sections[0].header.paragraphs[0].text='CLASSROOM DESK  /  DRAFT FOR REVIEW'
        doc.styles['Normal'].font.name='Calibri'
        doc.styles['Normal'].font.size=DocPt(11)
        doc.add_heading(draft.title,0)
        doc.add_paragraph(draft.summary)
        for section in draft.sections[:60]:
            doc.add_heading(section.heading,level=1)
            for text in section.paragraphs: doc.add_paragraph(text)
            for text in section.bullets: doc.add_paragraph(text,style='List Bullet')
            if section.table_headers:
                table=doc.add_table(rows=1,cols=len(section.table_headers)); table.style='Light Shading Accent 1'
                for cell,text in zip(table.rows[0].cells,section.table_headers): cell.text=text
                for row in section.table_rows[:200]:
                    for cell,text in zip(table.add_row().cells,row): cell.text=text
            if section.source_refs:
                note=doc.add_paragraph('Sources: '+'; '.join(section.source_refs),style='Caption')
                for run in note.runs: run.italic=True
        if images:
            doc.add_heading('Images supplied with the assignment',level=1)
            doc.add_paragraph('These source images are included for reference. Review the assignment requirements and place them where needed.')
            for asset in images:
                doc.add_paragraph(asset['filename'])
                doc.add_picture(io.BytesIO(asset['data']),width=DocInches(5.7))
        for title,items in [('Sources supplied',draft.sources),('Assumptions',draft.assumptions),('Arithmetic checks',calculations),('Review notes',draft.review_notes)]:
            if items:
                doc.add_heading(title,level=1)
                for item in items: doc.add_paragraph(item,style='List Bullet')
        path=folder/'assignment.docx'; doc.save(path); files.append({'name':path.name,'kind':'DOCX'})
    if 'pptx' in formats:
        if not draft.slides: raise IntegrationError('The model returned no slides. Retry generation.')
        pres=Presentation(); pres.slide_width=Inches(13.333); pres.slide_height=Inches(7.5)
        def add_slide(title,bullets,notes='',source_refs=None):
            slide=pres.slides.add_slide(pres.slide_layouts[6])
            slide.background.fill.solid(); slide.background.fill.fore_color.rgb=RGBColor.from_string('101F20')
            bar=slide.shapes.add_shape(1, Inches(.6), Inches(.65), Inches(.08), Inches(.55)); bar.fill.solid(); bar.fill.fore_color.rgb=RGBColor.from_string('C2E49A'); bar.line.fill.background()
            heading=slide.shapes.add_textbox(Inches(.95),Inches(.55),Inches(11.5),Inches(1.1)).text_frame
            heading.word_wrap=True; heading.paragraphs[0].text=title
            heading.paragraphs[0].font.size=Pt(30); heading.paragraphs[0].font.bold=True; heading.paragraphs[0].font.color.rgb=RGBColor.from_string('F5F6F0')
            body=slide.shapes.add_textbox(Inches(.95),Inches(1.9),Inches(11.3),Inches(4.5)).text_frame; body.word_wrap=True
            for i,text in enumerate(bullets):
                p=body.paragraphs[0] if i==0 else body.add_paragraph(); p.text='• '+text; p.font.size=Pt(22 if len(bullets)<=5 else 18); p.font.color.rgb=RGBColor.from_string('DFE7DF'); p.space_after=Pt(18)
            foot=slide.shapes.add_textbox(Inches(.95),Inches(6.95),Inches(11),Inches(.25)).text_frame
            foot.paragraphs[0].text=f'DRAFT FOR REVIEW  •  {len(pres.slides):02d}'; foot.paragraphs[0].font.size=Pt(9); foot.paragraphs[0].font.color.rgb=RGBColor.from_string('AABBAA')
            if source_refs: notes=(notes+'\n\nSources: '+'; '.join(source_refs)).strip()
            slide.notes_slide.notes_text_frame.text=notes
        add_slide(draft.title,[draft.summary])
        for slide in draft.slides[:40]: add_slide(slide.title,slide.bullets,slide.notes,slide.source_refs)
        for asset in images:
            slide=pres.slides.add_slide(pres.slide_layouts[6])
            slide.background.fill.solid(); slide.background.fill.fore_color.rgb=RGBColor.from_string('101F20')
            bar=slide.shapes.add_shape(1, Inches(.6), Inches(.65), Inches(.08), Inches(.55)); bar.fill.solid(); bar.fill.fore_color.rgb=RGBColor.from_string('C2E49A'); bar.line.fill.background()
            heading=slide.shapes.add_textbox(Inches(.95),Inches(.55),Inches(11.5),Inches(.8)).text_frame
            heading.paragraphs[0].text='Source image: '+asset['filename']; heading.paragraphs[0].font.size=Pt(25); heading.paragraphs[0].font.bold=True; heading.paragraphs[0].font.color.rgb=RGBColor.from_string('F5F6F0')
            width,height=asset['size']; scale=min(11.3/width,5.25/height)
            image_width,image_height=width*scale,height*scale
            left=(13.333-image_width)/2; top=1.55+(5.25-image_height)/2
            slide.shapes.add_picture(io.BytesIO(asset['data']),Inches(left),Inches(top),width=Inches(image_width),height=Inches(image_height))
            footer=slide.shapes.add_textbox(Inches(.95),Inches(6.95),Inches(11),Inches(.25)).text_frame
            footer.paragraphs[0].text='SUPPLIED ASSIGNMENT IMAGE  •  REVIEW BEFORE SUBMITTING'; footer.paragraphs[0].font.size=Pt(9); footer.paragraphs[0].font.color.rgb=RGBColor.from_string('AABBAA')
        for title,items in [('Sources',draft.sources),('Assumptions and review',draft.assumptions+draft.review_notes),('Arithmetic checks',calculations)]:
            for start in range(0,len(items),5): add_slide(title,items[start:start+5])
        path=folder/'assignment.pptx'; pres.save(path); files.append({'name':path.name,'kind':'PPTX'})
    if 'xlsx' in formats: files.append(_make_workbook(draft,folder))
    return files
